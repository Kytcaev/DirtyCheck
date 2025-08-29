import os
import importlib
import datetime
from openpyxl import load_workbook
from tkinter import Tk, filedialog
import sys

def resource_path(relative_path):
    """ Get absolute path to resource, works for dev and for PyInstaller """
    if hasattr(sys, '_MEIPASS'):
        return os.path.join(sys._MEIPASS, relative_path)
    return os.path.join(os.path.abspath("."), relative_path)


TEMPLATE_FILE = resource_path("Summary of comments.xlsx")
OUTPUT_FILE = f"Summary_of_comments_{datetime.datetime.now().strftime('%Y%m%d')}.xlsx"

# Модули разделены на проверки уровня файла и папки
FILE_LEVEL_CHECKS = [
    "checks.check_cyrillic_filename",
    "checks.check_cyrillic_pdf",
    "checks.check_footer",
    "checks.check_comments",
    "checks.check_duplicates",
    "checks.check_standards",   # сюда подключен модуль стандартов
]

FOLDER_LEVEL_CHECKS = [
    "checks.check_package",
]


def find_first_empty_row(ws):
    """Ищем первую пустую строку в колонке B (код документа)"""
    row = 2
    while ws.cell(row=row, column=2).value not in (None, ""):
        row += 1
    return row


def main():
    # выбираем папку
    Tk().withdraw()
    folder = filedialog.askdirectory(title="Выберите папку с PDF")
    if not folder:
        print("Папка не выбрана.")
        return

    # загружаем шаблон
    wb = load_workbook(TEMPLATE_FILE)
    ws = wb.active

    # загружаем проверки
    file_checks = [importlib.import_module(m) for m in FILE_LEVEL_CHECKS]
    folder_checks = [importlib.import_module(m) for m in FOLDER_LEVEL_CHECKS]

    # собираем pdf
    files = []
    for root, _, filenames in os.walk(folder):
        for f in filenames:
            if f.startswith("PKS2") and f.lower().endswith(".pdf"):
                files.append(os.path.join(root, f))

    print(f"Найдено {len(files)} файлов для проверки.")

    # нумерация
    next_row = find_first_empty_row(ws)
    index = 1

    # --- Проверки для каждого файла ---
    for file_path in files:
        for check_module in file_checks:
            try:
                results = check_module.check(file_path)
            except Exception as e:
                print(f"[Ошибка] {check_module.__name__}: {e}")
                results = []

            if not results:
                continue

            for res in results:
                if len(res) == 3:
                    code, comment, author = res
                    page = None
                else:
                    code, page, comment, author = res

                ws.cell(row=next_row, column=1).value = index
                ws.cell(row=next_row, column=2).value = code
                ws.cell(row=next_row, column=3).value = comment
                ws.cell(row=next_row, column=4).value = author
                if page:
                    ws.cell(row=next_row, column=5).value = page

                next_row += 1
                index += 1

    # --- Проверки для всей папки ---
    for check_module in folder_checks:
        try:
            results = check_module.check(folder, files)
        except Exception as e:
            print(f"[Ошибка] {check_module.__name__}: {e}")
            results = []

        for res in results:
            code, comment, author = res
            ws.cell(row=next_row, column=1).value = index
            ws.cell(row=next_row, column=2).value = code
            ws.cell(row=next_row, column=3).value = comment
            ws.cell(row=next_row, column=4).value = author
            next_row += 1
            index += 1

    # сохраняем основной отчёт
    out_path = os.path.join(folder, OUTPUT_FILE)
    wb.save(out_path)
    print(f"\nГотово! Результаты сохранены в {out_path}")

    # --- отдельный отчёт по стандартам ---
    try:
        from checks import check_standards
        check_standards.finalize(folder)
    except Exception as e:
        print(f"[Стандарты] не удалось сохранить отдельный отчёт: {e}")


if __name__ == "__main__":
    main()
    input("\nГотово. Нажмите Enter, чтобы закрыть окно...")
